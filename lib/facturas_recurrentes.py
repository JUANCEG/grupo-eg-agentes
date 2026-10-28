#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""facturas_recurrentes.py — alerta de facturas RECURRENTES que no se han emitido (SIIGO, las 4 empresas del Grupo EG).

Origen: en septiembre-2026 Tecnolavex no facturo a un cliente mensual de valor casi fijo (~$22,86 M) y nadie lo noto hasta octubre.

Que hace
  1. Descarga de SIIGO las facturas de ~14 meses de cada empresa (solo lectura).
  2. Un cliente es RECURRENTE si tiene factura en >= 4 de los 6 meses completos anteriores (o en los 3 ultimos meses seguidos).
  3. Para el mes en curso marca: FALTA (sin factura), VALOR DISTINTO (p. ej. 3 meses facturados juntos), y revisa si en los 3 meses anteriores se SALTO alguno.
  4. Calcula el valor habitual (mediana) y si es FIJO (variacion <= 5 %) o variable, y el dia habitual de emision.
Modos: --modo dia25 (solo corre el dia 25), --modo ultimo-habil (solo el ultimo dia habil del mes, calendario de festivos de Colombia), --modo ahora (cualquier dia).
Credenciales SOLO por variables de entorno SIIGO_<EMPRESA>_USER / SIIGO_<EMPRESA>_KEY (nunca se imprimen).
"""
import argparse
import json
import os
import statistics
import sys
import time
import urllib.request
from calendar import monthrange
from collections import defaultdict
from datetime import date, datetime, timedelta

EMPRESAS = ("omicron", "klar", "lavaclinicas", "tecnolavex")
NOMBRE = {"omicron": "Omicron", "klar": "Klar Solutions", "lavaclinicas": "Lavaclinicas", "tecnolavex": "Tecnolavex"}


# ----------------------------- calendario de Colombia -----------------------------
def pascua(y):
    a, b, c = y % 19, y // 100, y % 100
    d, e = b // 4, b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = c // 4, c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    mes, dia = divmod(h + l - 7 * m + 114, 31)
    return date(y, mes, dia + 1)


def _lunes(d):
    """Ley Emiliani: el festivo se traslada al lunes siguiente (si ya es lunes, se queda)."""
    return d if d.weekday() == 0 else d + timedelta(days=7 - d.weekday())


def festivos_colombia(y):
    e = pascua(y)
    fijos = [date(y, 1, 1), date(y, 5, 1), date(y, 7, 20), date(y, 8, 7), date(y, 12, 8), date(y, 12, 25), e - timedelta(days=3), e - timedelta(days=2)]
    trasladados = [date(y, 1, 6), date(y, 3, 19), date(y, 6, 29), date(y, 8, 15), date(y, 10, 12), date(y, 11, 1), date(y, 11, 11),
                   e + timedelta(days=39), e + timedelta(days=60), e + timedelta(days=68)]
    return set(fijos) | {_lunes(d) for d in trasladados}


def es_habil(d):
    return d.weekday() < 5 and d not in festivos_colombia(d.year)


def ultimo_habil(y, m):
    d = date(y, m, monthrange(y, m)[1])
    while not es_habil(d):
        d -= timedelta(days=1)
    return d


def fechas_de_control(desde, meses=13):
    """Lista de (fecha, tipo) en que DEBE correr la revision: el dia 25 de cada mes y el ultimo dia habil."""
    out, y, m = [], desde.year, desde.month
    for _ in range(meses):
        for d, t in ((date(y, m, 25), "dia25"), (ultimo_habil(y, m), "ultimo-habil")):
            if d >= desde:
                out.append((d, t))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return sorted(out)


# ----------------------------- analisis (funciones puras, con pruebas) -----------------------------
def mes_de(s):
    return s[:7]


def meses_anteriores(hoy, n):
    y, m, out = hoy.year, hoy.month, []
    for _ in range(n):
        y, m = (y - 1, 12) if m == 1 else (y, m - 1)
        out.append(f"{y}-{m:02d}")
    return out


def analizar(facturas, hoy, *, ventana=6, minimo_meses=4):
    """facturas: lista de dicts {cliente, fecha:'AAAA-MM-DD', total, numero}. Devuelve lista de hallazgos por cliente recurrente."""
    cur = f"{hoy.year}-{hoy.month:02d}"
    prev = meses_anteriores(hoy, ventana)               # prev[0] = mes pasado
    por = defaultdict(lambda: defaultdict(list))
    for f in facturas:
        if f.get("cliente") and f.get("fecha"):
            por[f["cliente"]][mes_de(f["fecha"])].append(f)
    hallazgos = []
    for cli, meses in por.items():
        con = [m for m in prev if m in meses]
        ultimos3_seguidos = all(m in meses for m in prev[:3])
        if len(con) < minimo_meses and not ultimos3_seguidos:
            continue
        totales = [sum(x["total"] or 0 for x in meses[m]) for m in con]
        mediana = statistics.median(totales)
        media = statistics.mean(totales)
        fijo = len(totales) >= 3 and media > 0 and statistics.pstdev(totales) / media <= 0.05
        dias = [int(min(x["fecha"] for x in meses[m])[8:10]) for m in con]
        dia_habitual = int(statistics.median(dias))
        todas = sorted((x for ms in meses.values() for x in ms), key=lambda x: x["fecha"])
        ult = todas[-1]
        facturado_mes = cur in meses
        total_mes = sum(x["total"] or 0 for x in meses[cur]) if facturado_mes else 0
        saltados = [m for m in prev[:3] if m not in meses]
        estado = []
        if not facturado_mes:
            estado.append("FALTA")
        elif mediana and ((fijo and abs(total_mes - mediana) / mediana > 0.15) or (not fijo and (total_mes > 2 * mediana or total_mes < 0.4 * mediana))):
            estado.append("VALOR_DISTINTO")
        if saltados:
            estado.append("MES_SALTADO")
        if not facturado_mes and prev[0] not in meses and prev[1] not in meses:
            estado.append("POSIBLE_BAJA")
        if not estado:
            continue
        mes_pasado_saltado = prev[0] in saltados      # lo que paso con Tecnolavex: el mes anterior quedo sin factura
        if "POSIBLE_BAJA" not in estado and (("FALTA" in estado and (fijo or dia_habitual <= hoy.day)) or mes_pasado_saltado):
            urgencia = "ALTA"
        else:
            urgencia = "MEDIA" if "FALTA" in estado or "MES_SALTADO" in estado else "BAJA"
        hallazgos.append({"cliente": cli, "estado": estado, "urgencia": urgencia, "meses_con_factura_de_6": len(con), "valor_habitual": round(mediana), "valor_fijo": fijo,
                          "dia_habitual": dia_habitual, "ya_deberia_estar": dia_habitual <= hoy.day, "meses_saltados": saltados, "mes_pasado_saltado": mes_pasado_saltado,
                          "total_mes_actual": round(total_mes), "ultima_factura": {"numero": ult.get("numero"), "fecha": ult["fecha"], "total": round(ult["total"] or 0)}})
    orden = {"ALTA": 0, "MEDIA": 1, "BAJA": 2}
    return sorted(hallazgos, key=lambda h: (orden[h["urgencia"]], -h["valor_habitual"]))


def debe_correr(modo, hoy):
    if modo == "ahora":
        return True, ""
    if modo == "dia25":
        return hoy.day == 25, "hoy no es 25"
    if modo == "ultimo-habil":
        return hoy == ultimo_habil(hoy.year, hoy.month), f"hoy no es el ultimo dia habil del mes ({ultimo_habil(hoy.year, hoy.month)})"
    raise ValueError(modo)


def moneda(v):
    return "$" + f"{v:,.0f}".replace(",", ".")


def texto_aviso(resultados, hoy, modo):
    """Aviso corto para el celular. resultados: {empresa: {'hallazgos': [...], 'error': str|None}}"""
    cuando = {"dia25": "dia 25", "ultimo-habil": "cierre de mes", "ahora": "revision"}[modo]
    faltan = {NOMBRE[e]: sum(1 for h in r["hallazgos"] if "FALTA" in h["estado"] and h["urgencia"] != "BAJA") for e, r in resultados.items()}
    errores = [NOMBRE[e] for e, r in resultados.items() if r.get("error")]
    tot = sum(faltan.values())
    pasado = sum(1 for r in resultados.values() for h in r["hallazgos"] if h.get("mes_pasado_saltado") and "POSIBLE_BAJA" not in h["estado"])
    if tot or pasado:
        det = ", ".join(f"{n} {k}" for k, n in faltan.items() if n)
        extra = f"; {pasado} con el mes pasado sin facturar" if pasado else ""
        return f"URGENTE facturas recurrentes ({cuando} {hoy:%d-%b}): faltan {tot}{' (' + det + ')' if det else ''}{extra}. Avisa a contabilidad hoy. Detalle en el Buzon."[:200]
    if errores:
        return f"URGENTE: no pude revisar SIIGO de {', '.join(errores)} ({cuando}). Revisa las facturas recurrentes a mano."[:190]
    return f"Facturas recurrentes ({cuando} {hoy:%d-%b}): todas al dia en las 4 empresas."


def texto_correo_contabilidad(resultados, hoy):
    """Borrador (NO se envia) para avisar a contabilidad."""
    lineas = [f"Asunto: URGENTE - facturas recurrentes sin emitir ({hoy:%d-%m-%Y})", "", "Hola, la revision de SIIGO muestra clientes recurrentes sin factura en el mes en curso:", ""]
    for e, r in resultados.items():
        faltan = [h for h in r["hallazgos"] if ("FALTA" in h["estado"] or h.get("mes_pasado_saltado")) and "POSIBLE_BAJA" not in h["estado"]]
        if faltan:
            lineas.append(f"{NOMBRE[e]}:")
            for h in faltan:
                que = "SIN factura este mes" if "FALTA" in h["estado"] else "mes pasado SIN factura (verificar que lo facturado este mes lo cubra)"
                lineas.append(f"  - {h.get('nombre') or h['cliente']} (NIT {h['cliente']}): {que}. Valor habitual {moneda(h['valor_habitual'])}{' (fijo)' if h['valor_fijo'] else ''}, "
                              f"se factura hacia el dia {h['dia_habitual']}; ultima factura {h['ultima_factura']['numero']} del {h['ultima_factura']['fecha']}.")
    lineas += ["", "Por favor confirmar si se emitira este mes o si hay una novedad (cambio de contrato, retraso, cliente dado de baja).", "Gracias."]
    return "\n".join(lineas)


# ----------------------------- SIIGO (solo lectura) -----------------------------
def _env(*names):
    miss = [n for n in names if not os.environ.get(n)]
    if miss:
        raise RuntimeError("faltan variables de entorno: " + ", ".join(miss))
    return [os.environ[n] for n in names]


def siigo_headers(emp):
    u, k = _env(f"SIIGO_{emp.upper()}_USER", f"SIIGO_{emp.upper()}_KEY")
    pid = os.environ.get("SIIGO_PARTNER_ID", "grupoEgAutomatizacion")
    req = urllib.request.Request("https://api.siigo.com/auth", data=json.dumps({"username": u, "access_key": k}).encode(), method="POST",
                                 headers={"Content-Type": "application/json", "Partner-Id": pid})
    tok = json.loads(urllib.request.urlopen(req, timeout=30).read().decode())["access_token"]
    return {"Authorization": f"Bearer {tok}", "Partner-Id": pid}


def _get(url, h):
    err = None
    for i in range(4):
        try:
            return json.loads(urllib.request.urlopen(urllib.request.Request(url, headers=h), timeout=40).read().decode())
        except Exception as e:  # noqa: BLE001
            err = e
            time.sleep(2 * (i + 1))
    raise RuntimeError(f"SIIGO no respondio: {type(err).__name__}")


def descargar_facturas(emp, hoy, meses=14, max_paginas=80):
    h = siigo_headers(emp)
    y, m = hoy.year, hoy.month
    for _ in range(meses):
        y, m = (y - 1, 12) if m == 1 else (y, m - 1)
    desde = date(y, m, 1).isoformat()
    out, p = [], 1
    while p <= max_paginas:
        res = (_get(f"https://api.siigo.com/v1/invoices?page={p}&page_size=100&created_start={desde}", h).get("results")) or []
        if not res:
            break
        for f in res:
            out.append({"cliente": str((f.get("customer") or {}).get("identification") or ""), "fecha": f.get("date"), "total": f.get("total") or 0, "numero": f.get("name")})
        p += 1
    return out, h


def nombre_cliente(ident, h, cache={}):
    if ident not in cache:
        try:
            c = (_get("https://api.siigo.com/v1/customers?identification=" + str(ident), h).get("results") or [{}])[0]
            cache[ident] = c.get("commercial_name") or " ".join(c.get("name") or []).strip() or "Cliente sin nombre"
        except Exception:  # noqa: BLE001
            cache[ident] = "Cliente " + str(ident)
    return cache[ident]


def revisar(empresas, hoy):
    res = {}
    for e in empresas:
        try:
            facturas, h = descargar_facturas(e, hoy)
            hall = analizar(facturas, hoy)
            for x in hall:
                x["nombre"] = nombre_cliente(x["cliente"], h)
            res[e] = {"hallazgos": hall, "error": None, "facturas_leidas": len(facturas)}
        except Exception as ex:  # noqa: BLE001
            res[e] = {"hallazgos": [], "error": f"{type(ex).__name__}: {str(ex)[:120]}", "facturas_leidas": 0}
    return res


def escribir_excel(res, hoy, ruta):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Facturas recurrentes"
    ws.append(["Empresa", "Urgencia", "Estado", "Cliente", "NIT", "Valor habitual", "Fijo?", "Dia habitual", "Meses con factura (de 6)", "Mes actual facturado", "Meses saltados", "Ultima factura", "Fecha", "Valor ultima"])
    for e, r in res.items():
        for h in r["hallazgos"]:
            ws.append([NOMBRE[e], h["urgencia"], ", ".join(h["estado"]), h.get("nombre", ""), h["cliente"], h["valor_habitual"], "si" if h["valor_fijo"] else "no", h["dia_habitual"],
                       h["meses_con_factura_de_6"], h["total_mes_actual"], ", ".join(h["meses_saltados"]), h["ultima_factura"]["numero"], h["ultima_factura"]["fecha"], h["ultima_factura"]["total"]])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="9C0006")
    for col, w in zip("ABCDEFGHIJKLMN", [16, 10, 30, 42, 14, 16, 8, 12, 14, 18, 22, 14, 12, 16]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
        if row[1].value == "ALTA":
            for c in row[:3]:
                c.fill = PatternFill("solid", fgColor="FFC7CE")
    ws.freeze_panes = "E2"
    r2 = wb.create_sheet("Resumen")
    r2.append(["Revision", hoy.isoformat()])
    for e, r in res.items():
        r2.append([NOMBRE[e], f"{len(r['hallazgos'])} hallazgos", f"{r['facturas_leidas']} facturas leidas", r["error"] or "ok"])
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    wb.save(ruta)
    return ruta


def main(argv=None):
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("--modo", choices=["dia25", "ultimo-habil", "ahora"], default="ahora")
    ap.add_argument("--hoy", help="AAAA-MM-DD (pruebas)")
    ap.add_argument("--empresas", nargs="*", default=list(EMPRESAS))
    ap.add_argument("--xlsx")
    ap.add_argument("--forzar", action="store_true")
    ap.add_argument("--fechas-control", action="store_true", help="imprime las fechas de control de los proximos 13 meses y sale")
    a = ap.parse_args(argv)
    hoy = date.fromisoformat(a.hoy) if a.hoy else date.today()
    if a.fechas_control:
        print(json.dumps([[d.isoformat(), t] for d, t in fechas_de_control(hoy)]))
        return 0
    ok, why = debe_correr(a.modo, hoy)
    if not ok and not a.forzar:
        print(json.dumps({"ejecutar": False, "motivo": why, "hoy": hoy.isoformat()}, ensure_ascii=False))
        return 0
    res = revisar(a.empresas, hoy)
    if a.xlsx:
        escribir_excel(res, hoy, a.xlsx)
    print(json.dumps({"ejecutar": True, "hoy": hoy.isoformat(), "modo": a.modo, "aviso": texto_aviso(res, hoy, a.modo), "correo_contabilidad": texto_correo_contabilidad(res, hoy),
                      "xlsx": a.xlsx, "resultados": res}, ensure_ascii=False, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())
